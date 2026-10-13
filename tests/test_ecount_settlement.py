# 2026-09-01 hoyeon.han: 이카운트 정산서 테스트 신규 작성
"""이카운트 판매현황내역 파서 + 정산서 파이프라인 연결 검증.

검증 앵커 (tests/이카운트_판매현황내역_20260901_정상.xlsx 실측):
- 유효행 4561 (원본 4564행 중 월소계/총합계/출력일시 3행 제외)
- 총합계 931,452,241 / 수량 5,094 (원본 푸터와 일치)
- 구분(주문유형): 발주 4332 / 변심반품 95 / 구매 41 / 로켓그로스 16
- 비고(셀러):     러브지나 5차 1334 / 우유커플 16차 479 / 에르마린 19차 372
- 거래처 30곳, 시온프로젝트 1070건 220,107,000원
"""

import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

# 2026-10-13 hoyeon.han: 이 파일만 checkout해도 conftest 없이 Src를 import한다.
_TESTS_DIR = Path(__file__).resolve().parent
_SRC_DIR = _TESTS_DIR.parent / "Src"
if str(_SRC_DIR) not in sys.path:
    sys.path.insert(0, str(_SRC_DIR))

import ecount_settlement as es
import seller_settlement as ss
import settlement as st_mod

# 2026-10-13 hoyeon.han: 개인정보 샘플은 저장소에 추가하지 않고, 없으면 skip한다.
ECOUNT_XLSX = _TESTS_DIR / "이카운트_판매현황내역_20260901_정상.xlsx"
ECOUNT_XLSX_SWAPPED = _TESTS_DIR / "이카운트_판매현황내역_20260901.xlsx"


@pytest.fixture(scope="module")
def ecount_parsed():
    """공용 conftest가 없어도 실제 샘플이 있는 환경에서만 앵커를 검증한다."""
    if not ECOUNT_XLSX.is_file():
        pytest.skip(f"테스트 원본 파일이 없습니다: {ECOUNT_XLSX}")
    return es.parse_ecount_xlsx(str(ECOUNT_XLSX))


@pytest.fixture(scope="module")
def ecount_df(ecount_parsed):
    """실제 샘플의 파싱 결과이며 합성 테스트는 이 fixture에 의존하지 않는다."""
    return ecount_parsed["df"]


@pytest.fixture(scope="module")
def ecount_parsed_swapped():
    """옛 양식 샘플도 선택적으로 읽고 파일이 없으면 해당 검증만 건너뛴다."""
    if not ECOUNT_XLSX_SWAPPED.is_file():
        pytest.skip(f"테스트 원본 파일이 없습니다: {ECOUNT_XLSX_SWAPPED}")
    return es.parse_ecount_xlsx(str(ECOUNT_XLSX_SWAPPED))

# ===== 파싱 기본 =====


def test_valid_row_count(ecount_parsed, ecount_df):
    """월소계/총합계/출력일시 3행이 빠지고 4561행만 남는다."""
    assert len(ecount_df) == 4561
    assert ecount_parsed["raw_rows"] == 4564
    assert ecount_parsed["dropped_rows"] == 3
    assert ecount_parsed["invalid_date_rows"] == 0


def test_header_and_company(ecount_parsed):
    """A1 '회사명 : (주)솔루미랩 / 2026/08/01 ~ 2026/09/01' 파싱."""
    assert ecount_parsed["company"] == "(주)솔루미랩"
    assert ecount_parsed["start_date"] == date(2026, 8, 1)
    assert ecount_parsed["end_date"] == date(2026, 9, 1)
    assert ecount_parsed["period_from_data"] is False
    # 기존 정산서와 동일한 제목/유형 계약
    assert ecount_parsed["title"] == es.TITLE_DEFAULT == "품목별판매현황"
    assert ecount_parsed["doc_type"] == st_mod.TYPE_SALE == "매출"
    assert ecount_parsed["sheet_name"] == "판매현황내역"


def test_parse_ecount_header_variants():
    """회사명에 '/' 가 있어도 기간을 먼저 찾아 안전하게 분리한다."""
    c, s, e = es.parse_ecount_header("회사명 : A/B상사 / 2026/08/01  ~ 2026/09/01 ")
    assert (c, s, e) == ("A/B상사", date(2026, 8, 1), date(2026, 9, 1))

    c, s, e = es.parse_ecount_header("회사명 : (주)솔루미랩")
    assert c == "(주)솔루미랩" and s is None and e is None


def test_footer_rows_excluded(ecount_df):
    """'2026/08  계' / '총합계' / 출력일시 행이 데이터에 남지 않는다."""
    assert "총합계" not in set(ecount_df["회사명"])
    assert not ecount_df["작성일자"].str.contains("계", na=False).any()


def test_merged_columns_dropped(ecount_df):
    """단가 F:H 병합으로 생기는 팬텀 컬럼이 제거되고 컬럼이 중복되지 않는다."""
    cols = list(ecount_df.columns)
    assert not any(str(c).startswith("Unnamed") for c in cols)
    assert len(cols) == len(set(cols)), f"중복 컬럼: {cols}"
    assert cols == es.PARSED_COLS


def test_nbsp_stripped(ecount_df):
    """빈 셀의 \xa0(non-breaking space)가 모두 제거된다."""
    for c in ["회사명", "품목", "구분", "비고"]:
        assert not ecount_df[c].str.contains("\xa0", regex=False).any()


def test_date_normalized(ecount_df):
    """20260803 → 2026-08-03 문자열로 정규화."""
    assert ecount_df["작성일자"].str.fullmatch(r"\d{4}-\d{2}-\d{2}").all()
    assert ecount_df["작성일자"].min() == "2026-08-03"
    assert ecount_df["작성일자"].max() == "2026-08-28"
    assert ecount_df["작성일자"].nunique() == 18


# ===== 검산 (이카운트 양식 변경 회귀 방지) =====


def test_total_matches_footer(ecount_parsed, ecount_df):
    assert int(ecount_df["합계"].sum()) == 931_452_241
    assert int(ecount_parsed["footer_total"]) == 931_452_241


def test_qty_matches_footer(ecount_parsed, ecount_df):
    assert int(ecount_df["수량"].sum()) == 5094
    assert int(ecount_parsed["footer_qty"]) == 5094


def test_verify_totals_reports_match(ecount_parsed, ecount_df):
    v = es.verify_totals(
        ecount_df, ecount_parsed["footer_total"], ecount_parsed["footer_qty"]
    )
    assert v["checked"] is True
    assert v["matches"] is True
    assert v["total_diff"] == 0
    assert v["qty_diff"] == 0


def test_verify_totals_without_footer(ecount_df):
    """푸터가 없으면 검산을 건너뛴다 (에러 아님)."""
    v = es.verify_totals(ecount_df, None)
    assert v["checked"] is False
    assert v["matches"] is None


# ===== 구분(주문유형) / 비고(셀러) =====


def test_order_type_column(ecount_df):
    """구분 = 주문유형. 발주/반품 등이 들어간다."""
    vc = ecount_df["구분"].value_counts()
    assert vc["발주"] == 4332
    assert vc["변심반품"] == 95
    assert vc["구매"] == 41
    assert vc["로켓그로스"] == 16
    assert int((ecount_df["구분"] == "").sum()) == 43


def test_seller_in_remark(ecount_df):
    """특이사항 → 비고. 셀러명이 들어가 시트 분할 기준이 된다."""
    vc = ecount_df["비고"].value_counts()
    assert vc["러브지나 5차"] == 1334
    assert vc["우유커플 16차"] == 479
    assert vc["에르마린 19차"] == 372
    assert int((ecount_df["비고"] == "").sum()) == 866


def test_extract_order_types(ecount_df):
    """주문유형 목록은 건수 많은 순, 빈 값 제외."""
    types = es.extract_order_types(ecount_df)
    assert types[0] == "발주"
    assert "" not in types
    assert len(types) == 19


def test_filter_by_order_types(ecount_df):
    assert len(es.filter_by_order_types(ecount_df, ["발주"])) == 4332
    assert len(es.filter_by_order_types(ecount_df, ["발주", "구매"])) == 4332 + 41
    # None = 전체 (필터 없음)
    assert len(es.filter_by_order_types(ecount_df, None)) == len(ecount_df)


def test_swap_not_detected_on_correct_file(ecount_parsed):
    """정상 형식 파일에서는 뒤바뀜 경고가 뜨지 않는다."""
    assert ecount_parsed["swap_warning"] is None


def test_swap_detected_on_swapped_file(ecount_parsed_swapped):
    """구분/특이사항이 뒤바뀐 옛 파일은 경고를 만들어 낸다 (자동 교정은 하지 않음)."""
    w = ecount_parsed_swapped["swap_warning"]
    assert w is not None
    assert w["score_remark"] > w["score_order_type"]
    assert w["score_remark"] >= es.SWAP_MIN_SCORE
    # 자동 교정하지 않으므로 값은 원본 그대로 (구분에 셀러명이 남아 있음)
    assert "러브지나 5차" in set(ecount_parsed_swapped["df"]["구분"])


# ===== 금액 / 특수 행 =====


def test_unit_price_raw_is_default(ecount_parsed, ecount_df):
    """기본은 원본 단가 그대로 (2026-09-01 사용자 결정).

    이카운트는 '합계 = 단가 + 배송비 + 도선료' 이고 배송비/도선료가 0이므로
    원본 단가는 합계와 같다.
    """
    assert ecount_parsed["unit_price_mode"] == es.UNIT_RAW
    assert (ecount_df["단가"] == ecount_df["합계"]).all()


def test_unit_price_recalc_mode():
    """recalc 모드는 합계 ÷ 수량. 수량 0 행은 0으로 둔다."""
    # from conftest import ECOUNT_XLSX
    # 2026-10-13 hoyeon.han: 위 직접 import 대신 파일 내 기본 샘플 경로로 폴백.

    if not ECOUNT_XLSX.exists():
        pytest.skip(f"테스트 원본 파일이 없습니다: {ECOUNT_XLSX}")
    p = es.parse_ecount_xlsx(str(ECOUNT_XLSX), unit_price_mode=es.UNIT_RECALC)
    df = p["df"]
    # 합계 컬럼은 어느 모드에서도 보존된다
    assert int(df["합계"].sum()) == 931_452_241
    assert (df.loc[df["수량"] == 0, "단가"] == 0).all()
    nz = df[df["수량"] != 0]
    assert (nz["합계"] - nz["단가"] * nz["수량"]).abs().max() <= 5


def test_negative_rows_preserved(ecount_df):
    """반품(음수) 행이 버려지지 않는다."""
    assert int((ecount_df["수량"] < 0).sum()) == 63
    assert int((ecount_df["합계"] < 0).sum()) == 64


def test_zero_qty_rows_preserved(ecount_df):
    """수량 0 이고 금액만 있는 행(수수료 등)도 살아남는다."""
    assert int((ecount_df["수량"] == 0).sum()) == 4


def test_fee_item_rows(ecount_df):
    """배송비/도선료는 컬럼이 아니라 품목명이 붙은 별도 행으로 들어온다."""
    assert int((ecount_df["품목"] == "배송비").sum()) == 90
    assert int((ecount_df["품목"] == "도선료").sum()) == 40


# ===== 거래처 =====


def test_vendor_extraction(ecount_df):
    """기존 extract_vendors 를 그대로 재사용할 수 있다."""
    vendors = st_mod.extract_vendors(ecount_df)
    assert len(vendors) == 30
    assert "시온프로젝트" in vendors
    assert "" not in vendors


def test_summarize_vendors(ecount_df):
    sv = es.summarize_vendors(ecount_df).set_index("거래처")
    assert int(sv.loc["㈜씨씨오와이", "건수"]) == 1880
    assert int(sv.loc["㈜씨씨오와이", "금액"]) == 384_671_400
    assert int(sv.loc["시온프로젝트", "건수"]) == 1070
    assert int(sv.loc["시온프로젝트", "금액"]) == 220_107_000


def test_stock_movement_vendor_present(ecount_df):
    """'재고(출고)'는 금액 0원 재고이동. 하드코딩 제외하지 않고 남겨 둔다."""
    sv = es.summarize_vendors(ecount_df).set_index("거래처")
    assert int(sv.loc["재고(출고)", "건수"]) == 26
    assert int(sv.loc["재고(출고)", "금액"]) == 0


# ===== 기존 파이프라인 연결 =====


def test_end_to_end_pipeline(ecount_parsed, ecount_df):
    """파싱 이후는 settlement.py 함수를 한 줄도 고치지 않고 재사용한다."""
    d = st_mod.filter_by_period(
        ecount_df, ecount_parsed["start_date"], ecount_parsed["end_date"]
    )
    assert len(d) == 4561

    d = st_mod.filter_by_vendor(d, "시온프로젝트")
    assert len(d) == 1070

    d = st_mod.drop_unused_columns(d)
    # 구분(주문유형)은 OUTPUT_COLS 에 없으므로 여기서 자동으로 떨어져 나간다
    assert list(d.columns) == list(st_mod.OUTPUT_COLS)
    assert "구분" not in d.columns

    sheets = st_mod.split_into_sheets(d)
    assert list(sheets.keys()) == [
        "전체",
        "일반",
        "11am 4차",
        "밍디씨 2차",
        "에르마린 18차",
        "에르마린 19차",
        "우유커플 16차",
        "우유커플이벤트",
        "황인선",
    ]
    assert len(sheets["전체"]) == 1070
    assert len(sheets["일반"]) == 27
    # 전체를 뺀 나머지 시트의 행 합이 전체와 같아야 누락이 없다
    assert sum(len(v) for k, v in sheets.items() if k != "전체") == 1070
    assert int(sheets["전체"]["합계"].sum()) == 220_107_000


def test_output_is_6_columns(ecount_parsed, ecount_df, tmp_path):
    """산출물은 기존 정산서와 동일한 6열. '구분'은 출력되지 않는다."""
    d = st_mod.drop_unused_columns(st_mod.filter_by_vendor(ecount_df, "시온프로젝트"))
    sheets = st_mod.split_into_sheets(d)
    out = st_mod.write_settlement_xlsx(
        out_path=str(tmp_path / "정산서.xlsx"),
        title=ecount_parsed["title"],
        period_line=st_mod.build_period_line(
            ecount_parsed["start_date"], ecount_parsed["end_date"], "시온프로젝트"
        ),
        sheets=sheets,
    )
    wb = openpyxl.load_workbook(out)
    ws = wb["전체"]
    assert ws.cell(row=1, column=1).value == "품목별판매현황"
    assert ws.cell(row=2, column=1).value == (
        "작성일자 : 2026-08-01 ~ 2026-09-01    거래처 : 시온프로젝트"
    )
    assert [ws.cell(row=3, column=c).value for c in range(1, 7)] == list(
        st_mod.OUTPUT_COLS
    )
    assert ws.cell(row=3, column=7).value is None  # 7번째 열 없음

    last = 4 + len(sheets["전체"])
    assert ws.cell(row=last, column=2).value == "1070건"
    assert ws.cell(row=last, column=3).value == "1,083"
    assert ws.cell(row=last, column=5).value == "220,107,000"
    # 전체 시트 금액 합계는 강조(빨강 글씨 / 노랑 채우기)
    tcell = ws.cell(row=last, column=5)
    assert tcell.font.color.rgb == "FFFF0000"
    assert tcell.fill.start_color.rgb == "FFFFFF00"


# ===== 에러 케이스 =====


def _write_xlsx(path, rows):
    """A1 + 머리글 + 데이터 행으로 최소 이카운트 형식 파일 생성."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = es.SHEET_NAME_DEFAULT
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    wb.save(path)
    return str(path)


_A1 = "회사명 : (주)솔루미랩 / 2026/08/01  ~ 2026/09/01 "
_HEADER = ["일자", "거래처명", "품목명", "구분", "수량", "단가", "합계", "특이사항"]


def test_missing_required_column_raises(tmp_path):
    path = _write_xlsx(
        tmp_path / "bad.xlsx",
        [[_A1], ["일자", "거래처명", "품목명"], ["20260803", "가", "나"]],
    )
    with pytest.raises(ValueError, match="필수 열이 없습니다"):
        es.parse_ecount_xlsx(path)


def test_no_valid_rows_raises(tmp_path):
    path = _write_xlsx(
        tmp_path / "empty.xlsx",
        [[_A1], _HEADER, ["총합계", "", "", "", 0, 0, 0, ""]],
    )
    with pytest.raises(ValueError, match="유효한 데이터 행이 없습니다"):
        es.parse_ecount_xlsx(path)


def test_duplicate_header_raises(tmp_path):
    path = _write_xlsx(
        tmp_path / "dup.xlsx",
        [
            [_A1],
            _HEADER + ["합계"],
            ["20260803", "가", "나", "발주", 1, 100, 100, "셀러", 100],
        ],
    )
    with pytest.raises(ValueError, match="중복된 열 이름"):
        es.parse_ecount_xlsx(path)


def test_too_few_rows_raises(tmp_path):
    path = _write_xlsx(tmp_path / "short.xlsx", [[_A1], _HEADER])
    with pytest.raises(ValueError, match="행 수가 부족"):
        es.parse_ecount_xlsx(path)


def test_period_fallback_from_data(tmp_path):
    """A1 에 기간이 없으면 데이터 일자 범위로 대체한다."""
    path = _write_xlsx(
        tmp_path / "noperiod.xlsx",
        [
            ["회사명 : (주)솔루미랩"],
            _HEADER,
            ["20260805", "가", "나", "발주", 1, 100, 100, "셀러A"],
            ["20260810", "가", "다", "발주", 2, 200, 200, "셀러A"],
        ],
    )
    p = es.parse_ecount_xlsx(path)
    assert p["period_from_data"] is True
    assert p["start_date"] == date(2026, 8, 5)
    assert p["end_date"] == date(2026, 8, 10)


def test_invalid_date_rows_dropped_not_raised(tmp_path):
    """불가능한 날짜 1건 때문에 전체가 막히지 않는다."""
    path = _write_xlsx(
        tmp_path / "baddate.xlsx",
        [
            [_A1],
            _HEADER,
            ["20260805", "가", "나", "발주", 1, 100, 100, "셀러A"],
            ["20260231", "가", "다", "발주", 1, 100, 100, "셀러A"],
        ],
    )
    p = es.parse_ecount_xlsx(path)
    assert p["invalid_date_rows"] == 1
    assert len(p["df"]) == 1


def test_unreadable_file_raises(tmp_path):
    path = tmp_path / "notexcel.xlsx"
    path.write_text("이건 엑셀이 아닙니다", encoding="utf-8")
    with pytest.raises(ValueError, match="파일을 열 수 없습니다"):
        es.parse_ecount_xlsx(str(path))


# 2026-10-13 hoyeon.han: 실제 샘플/conftest 없이 실행하는 정산서 어댑터 회귀검증.
_ORDER_SETTLEMENT_COLS = [
    "출고일", "업체명", "구분", "수하인명", "수하인주소", "제품",
    "연락처1", "연락처2", "특이사항", "수량", "상품매출", "판매배송비", "도선료",
]


def _make_ecount_df(*rows):
    """개인정보가 없는 합성 파싱 행을 만들고 지정한 값만 덮어쓴다."""
    base = {
        "작성일자": "2026-10-01",
        "회사명": "테스트업체",
        "품목": "상품A",
        "구분": "발주",
        "수량": 1,
        "단가": 30545.4545,
        "합계": 10000,
        "비고": "셀러A",
        "수하인명": "테스트수하인",
        "수하인주소": "테스트시 테스트구 검증로 0",
        "연락처1": "000-0000-0000",
        "연락처2": "000-0000-0001",
        "배송비": 0,
        "도선료": 0,
    }
    return pd.DataFrame([{**base, **row} for row in (rows or ({},))])


def test_parser_preserves_optional_fees_phone2_and_raw_unit_price(tmp_path):
    """원본 비용은 지우지 않고 숫자로, 연락처/원본 소수 단가는 그대로 보존한다."""
    path = _write_xlsx(
        tmp_path / "optional_columns.xlsx",
        [
            [_A1],
            _HEADER + ["배송비", "도선료", "연락처2"],
            ["20260805", "테스트업체", "상품A", "발주", 2,
             30545.4545, 33600, "셀러A", "1,500.5", "-75.25", "000-0000-0001"],
        ],
    )
    parsed = es.parse_ecount_xlsx(path)
    df = parsed["df"]
    assert parsed["unit_price_mode"] == es.UNIT_RAW
    assert list(df.columns) == es.PARSED_COLS
    assert df.iloc[0]["단가"] == pytest.approx(30545.4545, rel=0, abs=1e-8)
    assert df.iloc[0]["배송비"] == 1500.5
    assert df.iloc[0]["도선료"] == -75.25
    assert df.iloc[0]["연락처2"] == "000-0000-0001"
    # 파싱은 원본을 읽고, 출력 어댑터가 별도 품목 행 양식인지 검증한다.
    with pytest.raises(ValueError, match="배송비"):
        es.build_ecount_sheet_plan(df)


def test_parser_defaults_missing_optional_columns(tmp_path):
    """배송비/도선료/연락처2가 없는 최소 양식도 0/0/빈 문자열로 보충한다."""
    path = _write_xlsx(
        tmp_path / "missing_optional.xlsx",
        [[_A1], _HEADER,
         ["20260805", "테스트업체", "상품A", "발주", 1, 10000, 10000, "셀러A"]],
    )
    df = es.parse_ecount_xlsx(path)["df"]
    for col in ("배송비", "도선료"):
        assert pd.api.types.is_numeric_dtype(df[col])
        assert df[col].tolist() == [0]
    assert df["연락처2"].tolist() == [""]
    assert es.build_ecount_sheet_plan(df)["overall_df"]["합계"].sum() == 10000


def test_order_adapter_exact_columns_values_and_source_unchanged():
    """13열 순서, 날짜 타입, 구매자 정보, 합계 전액 대응 및 원본 불변을 검증한다."""
    df = _make_ecount_df({"수량": 2, "합계": 33600})
    before = df.copy(deep=True)
    orders = es.to_order_settlement_df(df)
    assert list(ss.COLUMN_MAP_SALE.values()) == _ORDER_SETTLEMENT_COLS
    assert list(orders.columns) == _ORDER_SETTLEMENT_COLS
    assert pd.api.types.is_datetime64_any_dtype(orders["출고일"])
    assert orders.iloc[0].tolist() == [
        pd.Timestamp("2026-10-01"), "테스트업체", "발주", "테스트수하인",
        "테스트시 테스트구 검증로 0", "상품A", "000-0000-0000", "000-0000-0001",
        "셀러A", 2, 33600, 0, 0,
    ]
    pd.testing.assert_frame_equal(df, before)


def test_order_adapter_missing_optional_columns_and_nondefault_index():
    """이전 파싱 결과의 선택 열 누락과 비연속 index 때문에 금액이 틀어지지 않는다."""
    df = _make_ecount_df({}, {"수량": 2, "합계": 20000})
    df = df.drop(columns=["배송비", "도선료", "연락처2"])
    df.index = [7, 42]
    orders = es.to_order_settlement_df(df)
    assert list(orders.columns) == _ORDER_SETTLEMENT_COLS
    assert orders["상품매출"].tolist() == [10000, 20000]
    assert orders["연락처2"].tolist() == ["", ""]
    assert orders[["판매배송비", "도선료"]].eq(0).all().all()
    assert orders["출고일"].notna().all()


@pytest.mark.parametrize("seller", ["", " \t ", "\xa0", None, float("nan"), pd.NA])
def test_order_adapter_blank_seller_is_general(seller):
    """빈 셀러는 일반으로 정규화되어 필터/시트 생성에서 탈락하지 않는다."""
    orders = es.to_order_settlement_df(_make_ecount_df({"비고": seller}))
    assert orders["특이사항"].tolist() == ["일반"]
    plan = es.build_ecount_sheet_plan(_make_ecount_df({"비고": seller}))
    assert list(plan["seller_sheets"]) == ["일반"]
    assert plan["detail_df"]["특이사항"].tolist() == ["일반"]
    assert plan["overall_df"]["합계"].sum() == 10000


@pytest.mark.parametrize("entrypoint", ["to_order_settlement_df", "build_ecount_sheet_plan"])
@pytest.mark.parametrize("fee_col", ["배송비", "도선료"])
@pytest.mark.parametrize("amount", [3300, -3300, 0.25, "3,300"])
def test_nonzero_source_fees_stop_settlement(entrypoint, fee_col, amount):
    """양수/음수/소수 비용을 추정하지 않고 해당 비용 열 안내와 함께 생성 중단한다."""
    df = _make_ecount_df({}, {fee_col: amount})
    before = df.copy(deep=True)
    with pytest.raises(ValueError, match=fee_col):
        getattr(es, entrypoint)(df)
    pd.testing.assert_frame_equal(df, before)


# 2026-10-13 hoyeon.han: 집계키는 원본 단가가 아닌 행 합계/수량과 셀러를 사용한다.
@pytest.mark.parametrize("column,value", [
    ("합계", 1000.5), ("합계", -1000.5), ("합계", float("inf")),
    ("수량", 1.5), ("수량", -1.5), ("수량", float("inf")),
])
def test_fractional_or_infinite_quantity_amount_stops_output(column, value):
    """2026-10-13 hoyeon.han: 원본 총계와 정수 출력이 달라지는 입력은 거부한다."""
    df = _make_ecount_df({column: value})
    before = df.copy(deep=True)
    with pytest.raises(ValueError, match=column):
        es.build_ecount_sheet_plan(df)
    pd.testing.assert_frame_equal(df, before)


def test_fractional_amount_groups_cannot_round_to_different_total():
    """2026-10-13 hoyeon.han: 원본 2001원이 그룹별 2000원으로 바뀌지 않게 차단."""
    df = _make_ecount_df({"품목": "상품A", "합계": 1000.5},
                         {"품목": "상품B", "합계": 1000.5})
    assert df["합계"].sum() == 2001
    with pytest.raises(ValueError, match="합계"):
        es.build_ecount_sheet_plan(df)


# 2026-10-13 hoyeon.han: 정상 정수 입력의 집계 및 소수 단가 계산은 계속 지원.
def test_summary_combines_same_keys_and_recalculates_price():
    df = _make_ecount_df(
        {"수량": 2, "합계": 20000, "단가": 30545.4545},
        {"수량": 3, "합계": 30000, "단가": 777},
    )
    before = df.copy(deep=True)
    plan = es.build_ecount_sheet_plan(df)
    summary = plan["seller_sheets"]["셀러A"]
    assert list(summary.columns) == list(st_mod.OUTPUT_COLS)
    assert summary.to_dict("records") == [{
        "작성일자": "2026-10-01", "품목": "상품A", "수량": 5,
        "단가": 10000, "합계": 50000, "비고": "셀러A",
    }]
    assert len(plan["detail_df"]) == 2
    assert plan["detail_df"]["상품매출"].tolist() == [20000, 30000]
    assert plan["overall_df"]["합계"].sum() == 50000
    pd.testing.assert_frame_equal(df, before)


@pytest.mark.parametrize(
    "changed",
    [
        pytest.param({"합계": 12000}, id="different-price"),
        pytest.param({"비고": "셀러B"}, id="different-seller"),
        pytest.param({"수량": -1, "합계": -10000}, id="return"),
        pytest.param({"작성일자": "2026-10-02"}, id="different-date"),
        pytest.param({"품목": "상품B"}, id="different-item"),
    ],
)
def test_summary_keeps_each_group_key_separate(changed):
    df = _make_ecount_df({}, changed)
    plan = es.build_ecount_sheet_plan(
        df, remark_to_sheet={"셀러A": "통합", "셀러B": "통합"}
    )
    summary = plan["seller_sheets"]["통합"]
    assert len(summary) == 2
    for raw in df.to_dict("records"):
        expected_price = round(raw["합계"] / raw["수량"], 4)
        match = summary[
            (summary["작성일자"] == raw["작성일자"])
            & (summary["품목"] == raw["품목"])
            & (summary["단가"] == expected_price)
            & ((summary["수량"] < 0) == (raw["수량"] < 0))
            & (summary["비고"] == raw["비고"])
        ]
        assert len(match) == 1
        assert match.iloc[0]["수량"] == raw["수량"]
        assert match.iloc[0]["합계"] == raw["합계"]
    assert summary["합계"].sum() == df["합계"].sum()


def test_zero_price_sale_and_return_do_not_cancel_quantity():
    """가격 0인 판매와 반품도 음수 여부 키가 있어 서로 상계되지 않는다."""
    df = _make_ecount_df({"합계": 0}, {"수량": -1, "합계": 0})
    summary = es.build_ecount_sheet_plan(df)["seller_sheets"]["셀러A"]
    assert len(summary) == 2
    assert sorted(summary["수량"].tolist()) == [-1, 1]
    assert summary["단가"].tolist() == [0, 0]
    assert summary["합계"].sum() == 0


def test_zero_quantity_amount_survives_summary_and_detail():
    df = _make_ecount_df({"품목": "수수료", "수량": 0, "합계": 5000})
    plan = es.build_ecount_sheet_plan(df)
    summary = plan["seller_sheets"]["셀러A"]
    assert summary[["수량", "단가", "합계"]].iloc[0].tolist() == [0, 0, 5000]
    assert plan["detail_df"][["수량", "상품매출"]].iloc[0].tolist() == [0, 5000]
    assert plan["overall_df"]["합계"].sum() == 5000


def test_parser_keeps_raw_price_but_summary_uses_total_divided_by_quantity(tmp_path):
    path = _write_xlsx(
        tmp_path / "raw_price.xlsx",
        [[_A1], _HEADER,
         ["20260805", "테스트업체", "상품A", "발주", 2, 30545.4545, 33600, "셀러A"]],
    )
    df = es.parse_ecount_xlsx(path)["df"]
    before = df.copy(deep=True)
    plan = es.build_ecount_sheet_plan(df)
    row = plan["seller_sheets"]["셀러A"].iloc[0]
    assert df.iloc[0]["단가"] == pytest.approx(30545.4545, rel=0, abs=1e-8)
    assert row["단가"] == 16800
    assert row["수량"] == 2 and row["합계"] == 33600
    assert plan["detail_df"].iloc[0]["상품매출"] == 33600
    pd.testing.assert_frame_equal(df, before)


def test_fee_item_rows_are_included_once_without_extra_cost_lines():
    df = _make_ecount_df(
        {"수량": 2, "합계": 20000},
        {"품목": "배송비", "합계": 3000},
        {"품목": "도선료", "수량": 0, "합계": 2000},
    )
    orders = es.to_order_settlement_df(df)
    assert orders[["판매배송비", "도선료"]].eq(0).all().all()
    plan = es.build_ecount_sheet_plan(df)
    summary = plan["seller_sheets"]["셀러A"].set_index("품목")
    assert set(summary.index) == {"상품A", "배송비", "도선료"}
    assert len(summary) == 3
    assert summary.loc["배송비", "합계"] == 3000
    assert summary.loc["도선료", "합계"] == 2000
    assert summary["합계"].sum() == 25000
    assert plan["overall_df"]["합계"].sum() == 25000
    assert plan["detail_df"]["상품매출"].sum() == 25000
    assert len(plan["detail_df"]) == len(df)


def test_mixed_blank_sellers_are_kept_in_general_sheet():
    df = _make_ecount_df(
        {"비고": "", "수량": 1, "합계": 10000},
        {"비고": None, "수량": 2, "합계": 20000},
        {"비고": " \t ", "수량": 3, "합계": 30000},
        {"수량": 1, "합계": 7000},
    )
    plan = es.build_ecount_sheet_plan(df)
    assert set(plan["seller_sheets"]) == {"일반", "셀러A"}
    general = plan["seller_sheets"]["일반"]
    assert general["수량"].tolist() == [6]
    assert general["합계"].tolist() == [60000]
    assert plan["detail_df"]["특이사항"].tolist() == ["일반", "일반", "일반", "셀러A"]
    assert plan["overall_df"]["합계"].sum() == 67000


def test_detail_preserves_thirteen_columns_and_each_buyer_row():
    df = _make_ecount_df(
        {},
        {"구분": "변심반품", "수하인명": "반품검증수하인", "수량": -1, "합계": -10000,
         "수하인주소": "반품검증주소", "연락처1": "000-0000-0002", "연락처2": ""},
    )
    detail = es.build_ecount_sheet_plan(df)["detail_df"]
    assert list(detail.columns) == _ORDER_SETTLEMENT_COLS
    assert detail.to_numpy().tolist() == [
        ["2026-10-01", "테스트업체", "발주", "테스트수하인",
         "테스트시 테스트구 검증로 0", "상품A", "000-0000-0000", "000-0000-0001",
         "셀러A", 1, 10000, 0, 0],
        ["2026-10-01", "테스트업체", "변심반품", "반품검증수하인", "반품검증주소",
         "상품A", "000-0000-0002", "", "셀러A", -1, -10000, 0, 0],
    ]


# 2026-10-13 hoyeon.han: 품목 정리와 통합 배치를 모든 상세 모드/전체 옵션으로 검증.
@pytest.mark.parametrize("detail_mode", [ss.DETAIL_MODE_SINGLE, ss.DETAIL_MODE_PER_SELLER])
def test_item_mapping_is_identical_in_summary_overall_and_detail(detail_mode):
    df = _make_ecount_df(
        {"품목": "원본상품A"},
        {"품목": "원본상품B", "수량": 2, "합계": 20000},
        {"품목": "배송비", "합계": 3000},
    )
    before = df.copy(deep=True)
    plan = es.build_ecount_sheet_plan(
        df, item_mapping={"원본상품A": "정리상품", "원본상품B": "정리상품"},
        detail_mode=detail_mode,
    )
    for summary in (plan["seller_sheets"]["셀러A"], plan["overall_df"]):
        assert set(summary["품목"]) == {"정리상품", "배송비"}
        mapped = summary.loc[summary["품목"] == "정리상품"].iloc[0]
        assert mapped["수량"] == 3 and mapped["합계"] == 30000
        assert summary["합계"].sum() == 33000
    detail = (plan["detail_df"] if detail_mode == ss.DETAIL_MODE_SINGLE
              else plan["detail_sheets"]["상세내역_셀러A"])
    assert detail["제품"].tolist() == ["정리상품", "정리상품", "배송비"]
    assert detail["상품매출"].sum() == 33000
    pd.testing.assert_frame_equal(df, before)


@pytest.mark.parametrize("include_overall", [True, False])
@pytest.mark.parametrize(
    "detail_mode", [ss.DETAIL_MODE_NONE, ss.DETAIL_MODE_SINGLE, ss.DETAIL_MODE_PER_SELLER]
)
def test_merged_routing_detail_modes_and_existing_seller_plan_contract(
    tmp_path, detail_mode, include_overall
):
    df = _make_ecount_df(
        {}, {"비고": "셀러B", "수량": 2, "합계": 20000},
        {"비고": "", "합계": 4000},
    )
    routing = {"셀러A": "통합", "셀러B": "통합"}
    plan = es.build_ecount_sheet_plan(
        df, remark_to_sheet=routing, detail_mode=detail_mode, include_overall=include_overall
    )
    assert set(plan["seller_sheets"]) == {"통합", "일반"}
    merged = plan["seller_sheets"]["통합"]
    assert len(merged) == 2  # 시트만 합치며 집계키의 셀러는 유지한다.
    assert set(merged["비고"]) == {"셀러A", "셀러B"}
    assert merged["합계"].sum() == 30000
    assert sum(sdf["합계"].sum() for sdf in plan["seller_sheets"].values()) == 34000
    assert plan["detail_mode"] == detail_mode
    if include_overall:
        assert plan["overall_df"]["합계"].sum() == 34000
    else:
        assert plan["overall_df"] is None

    if detail_mode == ss.DETAIL_MODE_NONE:
        assert plan["detail_df"] is None and plan["detail_sheets"] is None
        detail_names = []
    elif detail_mode == ss.DETAIL_MODE_SINGLE:
        assert plan["detail_sheets"] is None
        assert len(plan["detail_df"]) == 3
        assert list(plan["detail_df"].columns) == _ORDER_SETTLEMENT_COLS
        assert plan["detail_df"]["상품매출"].sum() == 34000
        detail_names = [ss.DETAIL_SHEET_NAME]
    else:
        assert plan["detail_df"] is None
        assert set(plan["detail_sheets"]) == {
            "상세내역_셀러A", "상세내역_셀러B", "상세내역_일반"
        }
        for name, ddf in plan["detail_sheets"].items():
            assert list(ddf.columns) == _ORDER_SETTLEMENT_COLS
            seller = name.removeprefix("상세내역_")
            assert ddf["특이사항"].tolist() == [seller]
            assert ddf["상품매출"].sum() == {"셀러A": 10000, "셀러B": 20000, "일반": 4000}[seller]
        detail_names = list(plan["detail_sheets"])

    expected = ss.build_sheet_plan(
        es.to_order_settlement_df(df), ["셀러A", "셀러B", "일반"],
        seller_to_sheet=routing, detail_mode=detail_mode, include_overall=include_overall,
    )
    assert set(plan) == set(expected)
    for key in ("seller_sheets", "detail_sheets"):
        if expected[key] is None:
            assert plan[key] is None
        else:
            assert list(plan[key]) == list(expected[key])
            for name in expected[key]:
                pd.testing.assert_frame_equal(plan[key][name], expected[key][name])
    for key in ("detail_df", "overall_df"):
        if expected[key] is None:
            assert plan[key] is None
        else:
            pd.testing.assert_frame_equal(plan[key], expected[key])

    out = ss.write_seller_settlement_xlsx(
        str(tmp_path / "routing.xlsx"), ss.TITLE_SALE, "검증 기간", plan
    )
    wb = openpyxl.load_workbook(out)
    try:
        expected_names = ([ss.SHEET_OVERALL] if include_overall else [])
        expected_names += list(plan["seller_sheets"]) + detail_names
        assert wb.sheetnames == expected_names
        for name in detail_names:
            assert wb[name].max_column == 13
            assert [cell.value for cell in wb[name][1]] == _ORDER_SETTLEMENT_COLS
    finally:
        wb.close()


def test_sheet_plan_defaults_to_single_detail_and_overall():
    plan = es.build_ecount_sheet_plan(_make_ecount_df())
    assert plan["detail_mode"] == ss.DETAIL_MODE_SINGLE
    assert plan["detail_df"] is not None and plan["detail_sheets"] is None
    assert plan["overall_df"]["합계"].tolist() == [10000]


@pytest.mark.parametrize("name", [ss.DETAIL_SHEET_NAME, ss.SHEET_OVERALL])
@pytest.mark.parametrize("use_routing", [True, False])
def test_reserved_sheet_names_raise_for_routing_or_source_seller(name, use_routing):
    df = _make_ecount_df({"비고": "셀러A" if use_routing else name})
    routing = {"셀러A": name} if use_routing else None
    with pytest.raises(ValueError, match=name):
        es.build_ecount_sheet_plan(df, remark_to_sheet=routing)


def test_overall_name_is_allowed_when_overall_is_disabled():
    plan = es.build_ecount_sheet_plan(
        _make_ecount_df(), remark_to_sheet={"셀러A": ss.SHEET_OVERALL}, include_overall=False
    )
    assert list(plan["seller_sheets"]) == [ss.SHEET_OVERALL]
    assert plan["overall_df"] is None


def test_invalid_detail_mode_is_rejected():
    with pytest.raises(ValueError, match="상세내역 모드"):
        es.build_ecount_sheet_plan(_make_ecount_df(), detail_mode="unsupported")


# 2026-10-13 hoyeon.han: 실제 파일을 다시 열어 서식/금액/소수 단가 저장을 검증한다.
def test_seller_writer_format_totals_buyer_info_and_decimal_precision(tmp_path):
    df = _make_ecount_df(
        {"수량": 3, "합계": 100000},
        {"수량": -1, "합계": -33333, "구분": "변심반품"},
        {"품목": "배송비", "합계": 3000},
        {"품목": "수수료", "수량": 0, "합계": 5000},
    )
    plan = es.build_ecount_sheet_plan(df)
    period_line = ss.build_period_line(
        "테스트업체", start_date=date(2026, 10, 1), end_date=date(2026, 10, 1)
    )
    out = ss.write_seller_settlement_xlsx(
        str(tmp_path / "ecount_settlement.xlsx"), ss.TITLE_SALE, period_line, plan
    )
    wb = openpyxl.load_workbook(out)
    try:
        assert wb.sheetnames == [ss.SHEET_OVERALL, "셀러A", ss.DETAIL_SHEET_NAME]
        for name, summary in [
            (ss.SHEET_OVERALL, plan["overall_df"]), ("셀러A", plan["seller_sheets"]["셀러A"])
        ]:
            ws = wb[name]
            assert ws.max_column == 6
            assert ws.cell(1, 1).value == ss.TITLE_SALE
            assert ws.cell(2, 1).value == period_line
            assert {str(rng) for rng in ws.merged_cells.ranges} == {"A1:F1", "A2:F2"}
            assert [cell.value for cell in ws[3]] == list(st_mod.OUTPUT_COLS)
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        assert cell.font.size == 13
            for excel_row, expected in enumerate(summary.to_dict("records"), start=4):
                assert [ws.cell(excel_row, col).value for col in (1, 2, 6)] == [
                    expected["작성일자"], expected["품목"], expected["비고"]
                ]
                for col, field in ((3, "수량"), (5, "합계")):
                    assert ws.cell(excel_row, col).data_type == "n"
                    assert ws.cell(excel_row, col).value == expected[field]
                    assert ws.cell(excel_row, col).number_format == "#,##0"
                price = ws.cell(excel_row, 4)
                assert price.data_type == "n"
                assert price.value == pytest.approx(expected["단가"], rel=0, abs=1e-8)
                if expected["수량"] == 3:
                    assert price.value == pytest.approx(33333.3333, rel=0, abs=1e-8)
                    assert price.number_format == "#,##0.####"
                else:
                    assert price.number_format == "#,##0"

            total_row = 4 + len(summary)
            assert ws.cell(total_row, 2).value == "4건"
            assert ws.cell(total_row, 3).value == 3
            assert ws.cell(total_row, 5).value == 74667
            for col in range(1, 7):
                cell = ws.cell(total_row, col)
                assert cell.font.bold and cell.font.size == 13
                if col == 5:
                    assert cell.font.color.rgb == "FFFF0000"
                    assert cell.fill.patternType == "solid"
                    assert cell.fill.start_color.rgb == "FFFFFF00"
                else:
                    assert cell.fill.patternType is None
                    assert cell.font.color is None or cell.font.color.type != "rgb" or cell.font.color.rgb != "FFFF0000"

        detail = wb[ss.DETAIL_SHEET_NAME]
        assert detail.max_column == 13
        assert [cell.value for cell in detail[1]] == _ORDER_SETTLEMENT_COLS
        assert detail.freeze_panes == "A2"
        assert detail.cell(2, 4).value == "테스트수하인"
        assert detail.cell(2, 5).value == "테스트시 테스트구 검증로 0"
        assert detail.cell(2, 7).value == "000-0000-0000"
        assert detail.cell(2, 8).value == "000-0000-0001"
        assert detail.cell(3, 3).value == "변심반품"
        assert [detail.cell(row, 11).value for row in range(2, 6)] == [100000, -33333, 3000, 5000]
        for row in range(2, 6):
            assert detail.cell(row, 10).data_type == "n"
            assert detail.cell(row, 11).number_format == "#,##0"
            assert detail.cell(row, 12).value in (None, "")
            assert detail.cell(row, 13).value in (None, "")
        total_row = 2 + len(df)
        assert [detail.cell(total_row, col).value for col in (11, 12, 13)] == [74667, 0, 0]
        assert detail.cell(total_row + 3, 11).value == 74667
        assert detail.cell(total_row, 11).font.bold
        assert detail.cell(total_row + 3, 11).font.bold
        for row in (total_row + 1, total_row + 2):
            assert all(detail.cell(row, col).value is None for col in range(1, 14))
    finally:
        wb.close()


@pytest.mark.parametrize("reserved_name", [ss.DETAIL_SHEET_NAME, ss.SHEET_OVERALL])
def test_seller_writer_rejects_reserved_sheet_names(tmp_path, reserved_name):
    plan = es.build_ecount_sheet_plan(_make_ecount_df())
    bad_plan = {**plan, "seller_sheets": {reserved_name: plan["seller_sheets"]["셀러A"]}}
    out_path = tmp_path / "reserved.xlsx"
    with pytest.raises(ValueError, match=reserved_name):
        ss.write_seller_settlement_xlsx(str(out_path), ss.TITLE_SALE, "검증 기간", bad_plan)
    assert not out_path.exists()


def test_real_sample_uses_seller_pipeline_without_changing_total(ecount_df):
    """샘플이 있는 환경에서만 기존 검산 금액과 새 13열 상세/요약을 함께 확인한다."""
    df = st_mod.filter_by_vendor(ecount_df, "시온프로젝트")
    before = df.copy(deep=True)
    plan = es.build_ecount_sheet_plan(df)
    assert plan["overall_df"]["합계"].sum() == 220_107_000
    assert sum(sdf["합계"].sum() for sdf in plan["seller_sheets"].values()) == 220_107_000
    assert len(plan["detail_df"]) == 1070
    assert list(plan["detail_df"].columns) == _ORDER_SETTLEMENT_COLS
    assert plan["detail_df"]["상품매출"].sum() == 220_107_000
    assert "일반" in plan["seller_sheets"]
    assert int((plan["detail_df"]["특이사항"] == "일반").sum()) == 27
    assert plan["detail_df"][["판매배송비", "도선료"]].eq(0).all().all()
    pd.testing.assert_frame_equal(df, before)
